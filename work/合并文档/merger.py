from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path
from threading import Event
from typing import Callable, Iterable, Sequence

import xlrd
import xlwt


EXPECTED_HEADER = (
    "CODE",
    "DESCRIPTION",
    "ITEM #",
    "ITEM DESCRIPTION",
    "UPC",
    "MODEL #",
    "UNIT COST",
    "CLAIM QTY",
    "UOM",
    "CLAIM AMT",
)
OUTPUT_HEADER = ("Invoice number",) + EXPECTED_HEADER[:5] + ("MODEL\u00a0#",) + EXPECTED_HEADER[6:]
FOOTER_LABELS = (
    "TOTAL QUANTITY",
    "TOTAL FOR CLAIM #",
    "TOTAL FREIGHT FEE",
    "TOTAL HANDLING COST",
    "TOTAL DISPOSAL FEE",
)
OUTPUT_PATTERN = re.compile(r"^合并结果_\d{8}_\d{6}(?:_\d+)?\.xlsx?$", re.IGNORECASE)
INVOICE_PATTERN = re.compile(r"0{5}(\d+)$")
MAX_XLS_ROWS = 65_536


class InvalidWorkbook(ValueError):
    pass


class MergeCancelled(RuntimeError):
    pass


@dataclass(frozen=True)
class FileIssue:
    path: Path
    reason: str


@dataclass
class SourceData:
    path: Path
    rows: list[list[object]]


@dataclass
class MergeResult:
    output_path: Path
    total_files: int
    merged_files: int
    skipped_files: int
    data_rows: int
    issues: list[FileIssue] = field(default_factory=list)


@dataclass
class HtmlCell:
    text_parts: list[str]
    colspan: int = 1


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._stack: list[dict[str, object]] = []
        self.tables: list[list[list[HtmlCell]]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "table":
            self._stack.append({"rows": [], "row": None, "cell": None})
            return
        if not self._stack:
            return
        context = self._stack[-1]
        if tag == "tr":
            context["row"] = []
        elif tag in {"td", "th"} and context["row"] is not None:
            attr_map = {key.lower(): value for key, value in attrs}
            try:
                colspan = max(1, int(attr_map.get("colspan") or "1"))
            except ValueError:
                colspan = 1
            context["cell"] = HtmlCell([], colspan)
        elif tag == "br" and context["cell"] is not None:
            context["cell"].text_parts.append(" ")

    def handle_data(self, data: str) -> None:
        if self._stack and self._stack[-1]["cell"] is not None:
            self._stack[-1]["cell"].text_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if not self._stack:
            return
        context = self._stack[-1]
        if tag in {"td", "th"} and context["cell"] is not None:
            if context["row"] is not None:
                context["row"].append(context["cell"])
            context["cell"] = None
        elif tag == "tr" and context["row"] is not None:
            if context["row"]:
                context["rows"].append(context["row"])
            context["row"] = None
            context["cell"] = None
        elif tag == "table":
            finished = self._stack.pop()
            if finished["rows"]:
                self.tables.append(finished["rows"])


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\u00a0", " ").split())


def normalize_header(row: Sequence[object]) -> tuple[str, ...]:
    return tuple(normalize_text(value).upper() for value in row)


def _expand_html_row(row: Sequence[HtmlCell]) -> list[str]:
    values: list[str] = []
    for cell in row:
        values.append(normalize_text("".join(cell.text_parts)))
        values.extend([""] * (cell.colspan - 1))
    return values


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin1")


def _find_claim_table(text: str) -> list[list[str]]:
    parser = _TableParser()
    parser.feed(text)
    for table in parser.tables:
        expanded = [_expand_html_row(row) for row in table]
        for index, row in enumerate(expanded):
            if len(row) == 10 and normalize_header(row) == EXPECTED_HEADER:
                return expanded[index:]
    raise InvalidWorkbook("未找到固定的 10 列表头")


def _footer_start(rows: Sequence[Sequence[object]]) -> int:
    non_empty = [
        (index, row)
        for index, row in enumerate(rows)
        if any(normalize_text(value) for value in row)
    ]
    if len(non_empty) < len(FOOTER_LABELS):
        raise InvalidWorkbook("底部五项汇总不完整，需要手动处理")

    footer = non_empty[-len(FOOTER_LABELS) :]
    for (_, row), expected in zip(footer, FOOTER_LABELS):
        if len(row) != 10:
            raise InvalidWorkbook("表格不是固定的 10 列结构")
        label = normalize_text(row[5]).upper()
        matched = label.startswith(expected) if expected == "TOTAL FOR CLAIM #" else label == expected
        if not matched:
            raise InvalidWorkbook("底部五项汇总不完整，需要手动处理")
    return footer[0][0]


def _as_number(value: object) -> object:
    text = normalize_text(value)
    if not text:
        return ""
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return text
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _convert_row(row: Sequence[object]) -> list[object]:
    values = [normalize_text(value) for value in row]
    for index in (0, 2, 4, 6, 7, 9):
        values[index] = _as_number(values[index])
    values[5] = _as_number(values[5])
    return values


def _validate_and_clean(path: Path, rows: Sequence[Sequence[object]]) -> SourceData:
    if not rows or len(rows[0]) != 10 or normalize_header(rows[0]) != EXPECTED_HEADER:
        raise InvalidWorkbook("表头与固定的 10 列标准不一致")
    start = _footer_start(rows[1:])
    details = [row for row in rows[1 : 1 + start] if any(normalize_text(value) for value in row)]
    for row in details:
        if len(row) != 10:
            raise InvalidWorkbook("明细行不是固定的 10 列结构")
    return SourceData(path=path, rows=[_convert_row(row) for row in details])


def _read_html_xls(path: Path) -> SourceData:
    return _validate_and_clean(path, _find_claim_table(_read_text(path)))


def _read_binary_xls(path: Path) -> SourceData:
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise InvalidWorkbook(f"无法读取文件：{exc}") from exc
    try:
        if not book.nsheets:
            raise InvalidWorkbook("工作簿中没有工作表")
        sheet = book.sheet_by_index(0)
        rows: list[list[object]] = []
        for row_index in range(sheet.nrows):
            row = [sheet.cell_value(row_index, column) for column in range(min(sheet.ncols, 10))]
            row.extend([""] * (10 - len(row)))
            if sheet.ncols > 10 and any(
                normalize_text(sheet.cell_value(row_index, column))
                for column in range(10, sheet.ncols)
            ):
                raise InvalidWorkbook("检测到第 10 列之后仍有数据")
            rows.append(row)
        return _validate_and_clean(path, rows)
    finally:
        book.release_resources()


def read_source(path: Path) -> SourceData:
    prefix = path.read_bytes()[:8]
    if prefix == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _read_binary_xls(path)
    if b"<table" in path.read_bytes()[:4096].lower():
        return _read_html_xls(path)
    raise InvalidWorkbook("文件扩展名是 .xls，但内容不是可识别的 Excel 或 HTML 表格")


def is_generated_output(path: Path) -> bool:
    return bool(OUTPUT_PATTERN.match(path.name))


def invoice_number(path: Path) -> str:
    match = INVOICE_PATTERN.search(path.stem)
    if match is None:
        raise InvalidWorkbook("文件名中未找到连续五个 0 后接数字，无法提取 Invoice number")
    return match.group(1)


def discover_files(folder: Path) -> tuple[list[Path], list[Path]]:
    if not folder.is_dir():
        raise NotADirectoryError(f"文件夹不存在：{folder}")
    found = sorted(
        (
            path
            for path in folder.iterdir()
            if path.is_file() and path.suffix.lower() in {".xls", ".xlsx"} and not path.name.startswith("~$")
        ),
        key=lambda path: path.name.casefold(),
    )
    generated = [path for path in found if is_generated_output(path)]
    sources = [path for path in found if path.suffix.lower() == ".xls" and not is_generated_output(path)]
    return sources, generated


def _safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", name).strip() or "合并结果"
    return cleaned[:31]


def _styles() -> tuple[xlwt.XFStyle, xlwt.XFStyle]:
    font = xlwt.Font()
    font.name = "宋体"
    font.height = 220

    left_alignment = xlwt.Alignment()
    left_alignment.horz = xlwt.Alignment.HORZ_LEFT
    left_alignment.vert = xlwt.Alignment.VERT_TOP

    right_alignment = xlwt.Alignment()
    right_alignment.horz = xlwt.Alignment.HORZ_RIGHT
    right_alignment.vert = xlwt.Alignment.VERT_TOP

    left = xlwt.XFStyle()
    left.font = font
    left.alignment = left_alignment

    right = xlwt.XFStyle()
    right.font = font
    right.alignment = right_alignment
    return left, right


def _write_xls_output(path: Path, sheet_name: str, rows: Sequence[Sequence[object]], cancel: Event) -> None:
    left, right = _styles()
    workbook = xlwt.Workbook(encoding="utf-8", style_compression=2)
    sheet = workbook.add_sheet(_safe_sheet_name(sheet_name), cell_overwrite_ok=False)
    widths = (15.0, 4.75, 35.375, 9.75, 22.0, 12.0, 9.1, 10.3, 10.3, 3.75, 9.75)
    for index, width in enumerate(widths):
        sheet.col(index).width = int(width * 256)

    for column, value in enumerate(OUTPUT_HEADER):
        sheet.write(0, column, value, left if column < 3 else right)

    for row_index, row in enumerate(rows, start=1):
        if row_index % 100 == 0 and cancel.is_set():
            raise MergeCancelled("用户已取消本次合并")
        for column, value in enumerate(row):
            sheet.write(row_index, column, value, left if column < 3 else right)

    temp_path = path.with_name(path.stem + ".tmp.xls")
    try:
        workbook.save(str(temp_path))
        if cancel.is_set():
            raise MergeCancelled("用户已取消本次合并")
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _write_xlsx_output(path: Path, sheet_name: str, rows: Sequence[Sequence[object]], cancel: Event) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_sheet_name(sheet_name)
    widths = (15.0, 4.75, 35.375, 9.75, 22.0, 12.0, 9.1, 10.3, 10.3, 3.75, 9.75)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    font = Font(name="宋体", size=11)
    alignments = (Alignment(horizontal="left", vertical="top"),
                  Alignment(horizontal="right", vertical="top"))
    for row_index, values in enumerate((OUTPUT_HEADER, *rows), start=1):
        if row_index % 100 == 0 and cancel.is_set():
            raise MergeCancelled("用户已取消本次合并")
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = font
            cell.alignment = alignments[column >= 4]

    temp_path = path.with_name(path.stem + ".tmp.xlsx")
    try:
        workbook.save(str(temp_path))
        if cancel.is_set():
            raise MergeCancelled("用户已取消本次合并")
        os.replace(temp_path, path)
    finally:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink()


def _output_extension(output_format: str | None) -> str:
    selected = output_format or ("xlsx" if sys.platform == "darwin" else "xls")
    if selected not in {"xls", "xlsx"}:
        raise ValueError(f"不支持的输出格式：{selected}")
    return f".{selected}"


def next_output_path(folder: Path, now: datetime | None = None, *, output_format: str | None = None) -> Path:
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    extension = _output_extension(output_format)
    candidate = folder / f"合并结果_{stamp}{extension}"
    sequence = 2
    while candidate.exists():
        candidate = folder / f"合并结果_{stamp}_{sequence}{extension}"
        sequence += 1
    return candidate


ProgressCallback = Callable[[int, int, Path, bool, str], None]
LogCallback = Callable[[str], None]


def merge_files(
    files: Iterable[Path],
    output_folder: Path,
    *,
    cancel: Event | None = None,
    progress: ProgressCallback | None = None,
    log: LogCallback | None = None,
    output_path: Path | None = None,
    output_format: str | None = None,
) -> MergeResult:
    ordered = sorted((Path(path) for path in files), key=lambda path: path.name.casefold())
    cancel = cancel or Event()
    details: list[list[object]] = []
    issues: list[FileIssue] = []
    first_valid: Path | None = None

    for index, path in enumerate(ordered, start=1):
        if cancel.is_set():
            raise MergeCancelled("用户已取消本次合并")
        try:
            invoice = invoice_number(path)
            source = read_source(path)
            first_valid = first_valid or path
            details.extend([[invoice, *row] for row in source.rows])
            message = f"成功：{path.name}（{len(source.rows)} 条明细）"
            ok = True
        except Exception as exc:
            reason = str(exc) or exc.__class__.__name__
            issues.append(FileIssue(path, reason))
            message = f"跳过：{path.name}（{reason}）"
            ok = False
        if log:
            log(message)
        if progress:
            progress(index, len(ordered), path, ok, message)

    if cancel.is_set():
        raise MergeCancelled("用户已取消本次合并")
    if first_valid is None:
        raise InvalidWorkbook("没有可合并的有效 .xls 文件")
    if len(details) + 1 > MAX_XLS_ROWS:
        raise InvalidWorkbook(f"合并后共有 {len(details) + 1} 行，超过约定的 65,536 行上限")

    destination = output_path or next_output_path(output_folder, output_format=output_format)
    if destination.suffix.lower() == ".xlsx":
        _write_xlsx_output(destination, first_valid.stem, details, cancel)
    else:
        _write_xls_output(destination, first_valid.stem, details, cancel)
    return MergeResult(
        output_path=destination,
        total_files=len(ordered),
        merged_files=len(ordered) - len(issues),
        skipped_files=len(issues),
        data_rows=len(details),
        issues=issues,
    )

