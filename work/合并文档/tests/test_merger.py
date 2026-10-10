from __future__ import annotations

import sys
import tempfile
import tkinter as tk
import unittest
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import xlrd
from openpyxl import load_workbook
from tkinterdnd2 import TkinterDnD

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import MergeApp, main
from merger import MergeResult, discover_files, merge_files, read_source


HEADER = [
    "CODE",
    "DESCRIPTION",
    "ITEM #",
    "ITEM DESCRIPTION",
    "UPC",
    "MODEL&nbsp;#",
    "UNIT COST",
    "CLAIM QTY",
    "UOM",
    "CLAIM AMT",
]


def write_html_xls(
    path: Path,
    rows: list[list[str]],
    *,
    bad_header: bool = False,
    missing_footer: bool = False,
) -> None:
    header = HEADER.copy()
    if bad_header:
        header[-1] = "OTHER"
    body = ["<table><tbody>"]
    body.append("<tr>" + "".join(f"<td>{value}</td>" for value in header) + "</tr>")
    for row in rows:
        body.append("<tr>" + "".join(f"<td>{value}</td>" for value in row) + "</tr>")
    body.append("<tr><td colspan='10'>&nbsp;</td></tr>")
    footers = [
        ("TOTAL QUANTITY", "1"),
        ("TOTAL FOR CLAIM # 123", "-1.00"),
        ("TOTAL FREIGHT FEE", "0.00"),
        ("TOTAL HANDLING COST", "0.00"),
        ("TOTAL DISPOSAL FEE", "0.00"),
    ]
    if missing_footer:
        footers.pop()
    for label, value in footers:
        body.append(
            f"<tr><td colspan='5'>&nbsp;</td><td colspan='3'>{label}</td>"
            f"<td colspan='2'>{value}</td></tr>"
        )
    body.append("</tbody></table>")
    path.write_text("\n".join(body), encoding="utf-8")


class MergerTests(unittest.TestCase):
    def test_merges_in_filename_order_and_removes_footers(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            write_html_xls(
                folder / "claim_details_000005301102474.xls",
                [["94", "B", "2", "ITEM B", "0002", "M2", "-2.5", "2", "EA", "-5"]],
            )
            write_html_xls(
                folder / "claim_details_000000400699589.xls",
                [["94", "A", "1", "ITEM A", "0001", "M1", "-1", "1", "EA", "-1"]],
            )

            files, generated = discover_files(folder)
            self.assertEqual([], generated)
            result = merge_files(files, folder, output_format="xls")

            self.assertEqual(2, result.merged_files)
            self.assertEqual(2, result.data_rows)
            book = xlrd.open_workbook(str(result.output_path))
            sheet = book.sheet_by_index(0)
            self.assertEqual(3, sheet.nrows)
            self.assertEqual(11, sheet.ncols)
            self.assertEqual("Invoice number", sheet.cell_value(0, 0))
            self.assertEqual("0400699589", sheet.cell_value(1, 0))
            self.assertEqual("5301102474", sheet.cell_value(2, 0))
            self.assertEqual(["CODE", "DESCRIPTION", "ITEM #", "ITEM DESCRIPTION", "UPC"], sheet.row_values(0)[1:6])
            self.assertEqual([94, "A", 1, "ITEM A", 1, "M1", -1, 1, "EA", -1], sheet.row_values(1)[1:])
            self.assertEqual("B", sheet.cell_value(2, 2))
            self.assertNotIn(
                "TOTAL",
                " ".join(
                    str(sheet.cell_value(row, column))
                    for row in range(sheet.nrows)
                    for column in range(sheet.ncols)
                ),
            )

    def test_bad_header_and_incomplete_footer_are_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            valid = folder / "claim_details_000001111111111.xls"
            bad_header = folder / "claim_details_000002222222222.xls"
            bad_footer = folder / "claim_details_000003333333333.xls"
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            write_html_xls(valid, [row])
            write_html_xls(bad_header, [row], bad_header=True)
            write_html_xls(bad_footer, [row], missing_footer=True)

            result = merge_files([valid, bad_header, bad_footer], folder, output_format="xls")

            self.assertEqual(1, result.merged_files)
            self.assertEqual(2, result.skipped_files)
            self.assertTrue(any("表头" in issue.reason for issue in result.issues))
            self.assertTrue(any("五项汇总" in issue.reason for issue in result.issues))

    def test_invoice_number_fills_every_detail_row_from_its_source_filename(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            six = folder / "claim_details_000005301106261.xls"
            two = folder / "claim_details_000005301111972.xls"
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            write_html_xls(six, [row] * 6)
            write_html_xls(two, [row] * 2)

            result = merge_files([two, six], folder, output_format="xls")

            sheet = xlrd.open_workbook(str(result.output_path)).sheet_by_index(0)
            self.assertEqual(8, result.data_rows)
            self.assertEqual(["5301106261"] * 6 + ["5301111972"] * 2, sheet.col_values(0)[1:])

    def test_filename_without_five_zero_prefix_is_skipped_and_logged(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            valid = folder / "claim_details_000000400699589.xls"
            invalid = folder / "claim_details_000040400699589.xls"
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            write_html_xls(valid, [row])
            write_html_xls(invalid, [row])
            messages: list[str] = []

            result = merge_files([invalid, valid], folder, log=messages.append, output_format="xls")

            self.assertEqual(1, result.merged_files)
            self.assertEqual(1, result.skipped_files)
            self.assertEqual(invalid, result.issues[0].path)
            self.assertIn("五个 0", result.issues[0].reason)
            self.assertTrue(any("跳过" in message and invalid.name in message for message in messages))

    def test_generated_outputs_are_excluded(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            write_html_xls(folder / "source.xls", [row])
            (folder / "合并结果_20261009_120000.xls").write_bytes(b"old")
            (folder / "合并结果_20261009_130000.xlsx").write_bytes(b"old")

            files, generated = discover_files(folder)

            self.assertEqual(["source.xls"], [path.name for path in files])
            self.assertEqual(
                ["合并结果_20261009_120000.xls", "合并结果_20261009_130000.xlsx"],
                [path.name for path in generated],
            )

    def test_macos_writes_real_xlsx_without_changing_merge_data(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            first = folder / "claim_details_000000400699589.xls"
            second = folder / "claim_details_000005301102474.xls"
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            write_html_xls(first, [row, row])
            write_html_xls(second, [row])

            with patch("merger.sys.platform", "darwin"):
                result = merge_files([second, first], folder)

            self.assertEqual(".xlsx", result.output_path.suffix)
            book = load_workbook(result.output_path, read_only=True, data_only=True)
            try:
                sheet = book.active
                self.assertEqual((4, 11), (sheet.max_row, sheet.max_column))
                self.assertEqual("Invoice number", sheet.cell(1, 1).value)
                self.assertEqual(["0400699589", "0400699589", "5301102474"],
                                 [sheet.cell(index, 1).value for index in (2, 3, 4)])
                self.assertEqual([94, "A", 1, "ITEM", 1, "M", -1, 1, "EA", -1],
                                 [sheet.cell(2, col).value for col in range(2, 12)])
                self.assertFalse(any("TOTAL" in str(cell.value) for cells in sheet for cell in cells))
            finally:
                book.close()

            files, generated = discover_files(folder)
            self.assertEqual([first, second], files)
            self.assertEqual([result.output_path], generated)

    def test_reads_nested_table_shape(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "nested.xls"
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            write_html_xls(path, [row])
            inner = path.read_text(encoding="utf-8")
            path.write_text(f"<table><tr><td>{inner}</td></tr></table>", encoding="utf-8")

            source = read_source(path)

            self.assertEqual(1, len(source.rows))
            self.assertEqual("A", source.rows[0][1])


class AppInteractionTests(unittest.TestCase):
    def test_gui_self_test_constructs_window_and_writes_report(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            report = Path(temp) / "gui.json"
            with patch("app.sys.argv", ["app.py", "--gui-self-test", str(report)]):
                with patch("app.TkinterDnD.Tk") as make_root:
                    with patch("app.MergeApp") as make_app:
                        make_root.return_value.TkdndVersion = "2.9.2"
                        code = main()

            self.assertTrue(report.exists(), "GUI self-test must write its result")
            self.assertEqual(0, code)
            self.assertEqual({"ok": True, "tkdnd": "2.9.2"}, json.loads(report.read_text(encoding="utf-8")))
            make_app.assert_called_once_with(make_root.return_value)
            make_root.return_value.mainloop.assert_not_called()

    @unittest.skipUnless(sys.platform == "win32", "Windows icon behavior")
    def test_window_uses_a_logo_and_updated_controls(self) -> None:
        root = TkinterDnD.Tk()
        root.withdraw()
        try:
            with patch.object(root, "iconbitmap", wraps=root.iconbitmap) as set_icon:
                app = MergeApp(root)
            app._draw_drop_zone(SimpleNamespace(width=700))
            drop_text = [
                app.drop.itemcget(item, "text")
                for item in app.drop.find_all()
                if app.drop.type(item) == "text"
            ]

            icon_path = Path(set_icon.call_args.args[0])
            self.assertEqual("app_icon.ico", icon_path.name)
            self.assertTrue(icon_path.is_file())
            self.assertIn("拖拽多个.xls文件或文件夹", drop_text)
            self.assertEqual("打开输出文件夹", app.open_button.cget("text"))
        finally:
            root.destroy()

    def test_mac_window_loads_png_icon_without_windows_ico(self) -> None:
        root = TkinterDnD.Tk()
        root.withdraw()
        try:
            with patch("app.sys.platform", "darwin"):
                with patch.object(root, "iconbitmap") as windows_icon:
                    with patch.object(root, "iconphoto") as mac_icon:
                        MergeApp(root)

            windows_icon.assert_not_called()
            mac_icon.assert_called_once()
            image = mac_icon.call_args.args[1]
            self.assertEqual("a-logo.png", Path(image.cget("file")).name)
        finally:
            root.destroy()

    def test_mac_window_starts_if_tk_rejects_window_icon(self) -> None:
        root = TkinterDnD.Tk()
        root.withdraw()
        try:
            with patch("app.sys.platform", "darwin"):
                with patch.object(root, "iconphoto", side_effect=tk.TclError("unsupported")):
                    try:
                        app = MergeApp(root)
                    except tk.TclError:
                        self.fail("An unsupported Tk window icon must not prevent macOS startup")
            self.assertEqual("合并文档 V5", root.title())
            self.assertEqual("disabled", str(app.start_button.cget("state")))
        finally:
            root.destroy()

    def test_progress_card_reveals_only_after_files_are_found(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "claim_001.xls").write_bytes(b"sample")
            root = TkinterDnD.Tk()
            root.withdraw()
            try:
                app = MergeApp(root)
                callbacks: list[tuple[object, tuple[object, ...]]] = []
                root.after = Mock(side_effect=lambda _delay, fn, *args: callbacks.append((fn, args)))
                self.assertEqual("", app.progress.master.winfo_manager())

                app._set_folder(folder)
                self.assertEqual("pack", app.progress_host.winfo_manager())
                self.assertLess(float(app.progress.master.place_info()["relwidth"]), 1.0)

                while callbacks:
                    callback, args = callbacks.pop(0)
                    callback(*args)
                self.assertAlmostEqual(1.0, float(app.progress.master.place_info()["relwidth"]))

                (folder / "claim_001.xls").unlink()
                app._set_folder(folder)
                self.assertEqual("", app.progress_host.winfo_manager())
            finally:
                root.destroy()

    def test_dragged_old_output_is_excluded_before_merge(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "claim_details_000005301111972.xls"
            old_output = folder / "合并结果_20261010_100000.xlsx"
            source.write_bytes(b"sample")
            old_output.write_bytes(b"old")

            root = TkinterDnD.Tk()
            root.withdraw()
            try:
                app = MergeApp(root)
                root.after = Mock()
                with patch("app.messagebox.showwarning"):
                    app._on_drop(SimpleNamespace(data=f"{{{source}}} {{{old_output}}}"))

                self.assertEqual("0 / 1", app.count_var.get())
                self.assertIn("1 个历史合并结果", app.log_text.get("1.0", "end"))
                with patch("app.messagebox.askyesno", return_value=True) as confirm:
                    with patch("app.threading.Thread") as worker:
                        app._start_merge()
                confirm.assert_called_once()
                self.assertEqual([source], worker.call_args.kwargs["args"][0])
            finally:
                root.destroy()

    def test_dragging_multiple_xls_scans_only_dropped_files_until_start(self) -> None:
        with tempfile.TemporaryDirectory(prefix="xls drag ") as temp:
            folder = Path(temp)
            row = ["94", "A", "1", "ITEM", "0001", "M", "-1", "1", "EA", "-1"]
            first = folder / "claim_details_000000400699589.xls"
            second = folder / "claim_details_000005301102474.xls"
            unselected = folder / "claim_details_000005301106261.xls"
            for path in (first, second, unselected):
                write_html_xls(path, [row])

            root = TkinterDnD.Tk()
            root.withdraw()
            try:
                app = MergeApp(root)
                root.after = Mock()
                event = SimpleNamespace(data=f"{{{second}}} {{{first}}}")
                with patch("app.messagebox.showwarning"):
                    app._on_drop(event)

                self.assertFalse(app.busy)
                self.assertEqual("0 / 2", app.count_var.get())
                self.assertIn("已读取 2 个 Excel 文件", app.status_var.get())
                self.assertIn("发现 2 个可处理的 .xls 文件", app.log_text.get("1.0", "end"))
                self.assertEqual("normal", str(app.start_button.cget("state")))
                self.assertFalse(list(folder.glob("合并结果_*.xls")))

                with patch("app.threading.Thread") as worker:
                    app._start_merge()
                self.assertEqual([first, second], worker.call_args.kwargs["args"][0])
            finally:
                root.destroy()

    @unittest.skipUnless(sys.platform == "win32", "Windows Explorer behavior")
    def test_open_output_selects_the_generated_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "合并结果_20261010_100000.xls"
            output.write_bytes(b"sample")
            root = TkinterDnD.Tk()
            root.withdraw()
            try:
                app = MergeApp(root)
                root.after = Mock()
                app.events.put(("done", MergeResult(output, 1, 1, 0, 1)))
                app._poll_events()
                self.assertEqual("normal", str(app.open_button.cget("state")))
                with patch("platform_actions.subprocess.Popen") as launch:
                    app._open_output()
                launch.assert_called_once_with(["explorer.exe", "/select,", str(output)])
            finally:
                root.destroy()

    def test_mac_open_output_reveals_generated_file_in_finder(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "合并结果_20261010_100000.xls"
            output.write_bytes(b"sample")
            root = TkinterDnD.Tk()
            root.withdraw()
            try:
                with patch("app.sys.platform", "darwin"):
                    app = MergeApp(root)
                    app.output_path = output
                    with patch("platform_actions.subprocess.Popen") as launch:
                        app._open_output()
                launch.assert_called_once_with(["open", "-R", str(output)])
            finally:
                root.destroy()

    def test_selecting_folder_only_scans_until_start_is_clicked(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "claim_001.xls").write_bytes(b"sample")
            (folder / "claim_002.xls").write_bytes(b"sample")
            (folder / "合并结果_20261009_120000.xls").write_bytes(b"old")

            root = TkinterDnD.Tk()
            root.withdraw()
            try:
                app = MergeApp(root)
                root.after = Mock()
                app._start_merge = Mock()

                app._set_folder(folder)

                app._start_merge.assert_not_called()
                self.assertEqual("0 / 2", app.count_var.get())
                self.assertIn("已读取 2 个 Excel 文件", app.status_var.get())
                self.assertEqual("normal", str(app.start_button.cget("state")))
            finally:
                root.destroy()


if __name__ == "__main__":
    unittest.main()
